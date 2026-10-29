"""Turn any dropped file into pages of word boxes (or, for office files, cell grids).

    read_pages(path, engine="auto", force_ocr=False) -> [Page]
    read_grids(path)                                  -> [grid]   (.xlsx/.xlsm/.docx)

A Page is {"words": [{text,x0,x1,top,bottom}], "width": float, "text": str, "source": str}.
Routing is PER PAGE: a PDF page with a real text layer is read exactly (pdfplumber); a page
without one is rasterized and OCR'd. Everything downstream is engine-agnostic.

OCR engines, all local/offline, each `PIL.Image -> (words, text)`:
  * winocr    - Windows.Media.Ocr, built into Windows 10+. No install. In-memory.
  * tesseract - external Tesseract binary via pytesseract, when installed.
Every image is EXIF-transposed ONCE here (_upright): phone photos carry their rotation in
EXIF, and an engine fed the raw pixels reads them upside-down.
"""
from __future__ import annotations
import os
from typing import Callable, Dict, List, Optional, Tuple

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".heic", ".heif"}
GRID_EXTS = {".xlsx", ".xlsm", ".docx"}
TEXT_LAYER_MIN_CHARS = 40   # a PDF page with fewer extractable chars is treated as a scan
OCR_DPI = 300
TEXT_HEIGHT_PX, MIN_TEXT_HEIGHT_PX = 28, 20     # re-read enlarged when the median word is shorter
MAX_SKEW_DEG = 0.3                              # re-read levelled when the text leans more


class NoOcrEngine(RuntimeError):
    """A page needs OCR and no local engine is usable."""


def ext_of(path: str) -> str:
    return os.path.splitext(path)[1].lower()


# ---- OCR engines --------------------------------------------------------------------------

def _winocr(img) -> Tuple[List[Dict], str]:
    import asyncio
    from winrt.windows.media.ocr import OcrEngine
    from winrt.windows.graphics.imaging import SoftwareBitmap, BitmapPixelFormat, BitmapAlphaMode
    from winrt.windows.storage.streams import DataWriter

    engine = OcrEngine.try_create_from_user_profile_languages()
    if engine is None:
        raise NoOcrEngine("Windows has no OCR language installed (Settings > Time & Language > Language).")
    img = img.convert("RGBA")
    scale = 1.0
    limit = OcrEngine.max_image_dimension
    if max(img.size) > limit:
        scale = max(img.size) / limit
        img = img.resize((int(img.width / scale), int(img.height / scale)))
    writer = DataWriter()
    writer.write_bytes(img.tobytes())
    bitmap = SoftwareBitmap.create_copy_with_alpha_from_buffer(
        writer.detach_buffer(), BitmapPixelFormat.RGBA8, img.width, img.height, BitmapAlphaMode.PREMULTIPLIED)

    async def recognize():
        return await engine.recognize_async(bitmap)

    result = asyncio.run(recognize())
    words = []
    for n, line in enumerate(result.lines):
        for w in line.words:
            r = w.bounding_rect
            words.append({"text": w.text, "x0": r.x * scale, "x1": (r.x + r.width) * scale,
                          "top": r.y * scale, "bottom": (r.y + r.height) * scale, "line": n})
    return words, "\n".join(line.text for line in result.lines)


def _winocr_available() -> bool:
    try:
        from winrt.windows.media.ocr import OcrEngine
        return OcrEngine.try_create_from_user_profile_languages() is not None
    except Exception:
        return False


def _tesseract_cmd() -> Optional[str]:
    try:
        import pytesseract
    except Exception:
        return None
    from shutil import which
    candidates = [
        os.environ.get("OCR_TESSERACT"),
        which("tesseract"),
        os.path.join(os.environ.get("ProgramFiles", ""), "Tesseract-OCR", "tesseract.exe"),
        os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "Tesseract-OCR", "tesseract.exe"),
    ]
    for c in candidates:
        if c and os.path.isfile(c):
            pytesseract.pytesseract.tesseract_cmd = c
            return c
    return None


def _tesseract(img) -> Tuple[List[Dict], str]:
    import tempfile
    import pytesseract
    if not _tesseract_cmd():
        raise NoOcrEngine("Tesseract is not installed.")
    # pytesseract hands the page to the binary through a temp file; keep that inside a
    # directory that is purged even on a crash, so no statement page lingers in %TEMP%.
    with tempfile.TemporaryDirectory(prefix="ledgerocr_") as tmp:
        prev, tempfile.tempdir = tempfile.tempdir, tmp
        try:
            d = pytesseract.image_to_data(img.convert("RGB"), output_type=pytesseract.Output.DICT,
                                          config="--psm 6")
        finally:
            tempfile.tempdir = prev
    words, lines = [], {}
    for i, raw in enumerate(d["text"]):
        txt = (raw or "").strip()
        if not txt:
            continue
        left, top = d["left"][i], d["top"][i]
        key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
        words.append({"text": txt, "x0": left, "x1": left + d["width"][i],
                      "top": top, "bottom": top + d["height"][i], "line": key})
        lines.setdefault(key, []).append(txt)
    return words, "\n".join(" ".join(ws) for ws in lines.values())


ENGINES: Dict[str, Tuple[Callable[[], bool], Callable]] = {
    "winocr": (_winocr_available, _winocr),
    "tesseract": (lambda: _tesseract_cmd() is not None, _tesseract),
}


def available_engines() -> List[str]:
    return [name for name, (ok, _) in ENGINES.items() if ok()]


def _engine(name: str) -> Tuple[str, Callable]:
    if name == "auto":
        found = available_engines()
        if not found:
            raise NoOcrEngine("This is a scan or photo, and no local OCR engine was found "
                              "(Windows OCR or Tesseract).")
        name = found[0]
    if name not in ENGINES:
        raise ValueError(f"unknown engine '{name}' (expected auto, {', '.join(ENGINES)})")
    return name, ENGINES[name][1]


# ---- pages --------------------------------------------------------------------------------

def _upright(img):
    from PIL import ImageOps
    return ImageOps.exif_transpose(img)


def _open_image(path: str):
    from PIL import Image
    if ext_of(path) in (".heic", ".heif"):
        try:
            import pillow_heif
        except ImportError:
            raise ValueError("HEIC/HEIF photos need the optional 'pillow-heif' package. "
                             "Or convert the photo to JPG/PNG.")
        pillow_heif.register_heif_opener()
    return Image.open(path)


def _skew_and_text_height(words: List[Dict]) -> Tuple[float, float]:
    """(degrees to rotate the image counter-clockwise to level the text, median word height)."""
    import math
    from statistics import median
    lines: Dict = {}
    for w in words:
        lines.setdefault(w["line"], []).append(w)
    slopes = []
    for ws in lines.values():
        a, b = min(ws, key=lambda w: w["x0"]), max(ws, key=lambda w: w["x0"])
        run_x = (b["x0"] + b["x1"] - a["x0"] - a["x1"]) / 2
        if len(ws) >= 3 and run_x > 10 * (a["bottom"] - a["top"]):      # long lines only
            slopes.append(((b["top"] + b["bottom"]) - (a["top"] + a["bottom"])) / 2 / run_x)
    angle = math.degrees(math.atan(median(slopes))) if len(slopes) >= 3 else 0.0
    return angle, median(w["bottom"] - w["top"] for w in words)


def _ocr_page(img, engine: str, unit: float = 1.0) -> Dict:
    """OCR one page. Engines want level text about 30 px tall (a 300-dpi scan); a skewed or
    low-resolution page is measured from the first pass, corrected, and read once more."""
    name, run = _engine(engine)
    img = _upright(img)
    words, text = run(img)
    if words:
        angle, height = _skew_and_text_height(words)
        zoom = min(3.0, TEXT_HEIGHT_PX / height) if height < MIN_TEXT_HEIGHT_PX else 1.0
        if abs(angle) > MAX_SKEW_DEG or zoom > 1.0:
            from PIL import Image
            fixed = img.convert("RGB").rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor="white")
            fixed = fixed.resize((int(fixed.width * zoom), int(fixed.height * zoom)), Image.LANCZOS)
            words2, text2 = run(fixed)
            if len(words2) >= 0.8 * len(words):      # never trade a good read for a worse one
                img, words, text, unit = fixed, words2, text2, unit * zoom
    if unit != 1.0:     # rasterized PDF page: back to PDF points, so columns carry across pages
        for w in words:
            for k in ("x0", "x1", "top", "bottom"):
                w[k] /= unit
    # Reading order from geometry, not from the engine: Windows OCR returns a table column by
    # column, which scrambles the plain-text output. Rows top-to-bottom, words left-to-right.
    from .tables import _cluster_rows
    text = "\n".join(" ".join(w["text"] for w in row) for row in _cluster_rows(words)) if words else text
    return {"words": words, "width": img.width / unit, "text": text, "source": name}


def _join_touching(words: List[Dict]) -> List[Dict]:
    """Some PDFs draw a stray space glyph on top of a digit, and pdfplumber splits there
    ('233.10' -> '2' '33.10'). A real word space has width; words that touch are one word."""
    out: List[Dict] = []
    for w in words:
        p = out[-1] if out else None
        if p and abs(p["top"] - w["top"]) < 1 and abs(w["x0"] - p["x1"]) < 0.1:
            out[-1] = dict(p, text=p["text"] + w["text"], x1=w["x1"], bottom=max(p["bottom"], w["bottom"]))
        else:
            out.append(w)
    return out


def read_pages(path: str, engine: str = "auto", force_ocr: bool = False) -> List[Dict]:
    if ext_of(path) in IMAGE_EXTS:
        img = _open_image(path)
        frames = getattr(img, "n_frames", 1)     # multi-page TIFF
        pages = []
        for i in range(frames):
            img.seek(i)
            pages.append(_ocr_page(img.copy(), engine))
        return pages
    import pdfplumber
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = "" if force_ocr else (page.extract_text() or "")
            if len(text) >= TEXT_LAYER_MIN_CHARS:
                # Gap tolerance relative to type size: spreadsheet exports leave cells ~0.35 em
                # apart with no space glyph, which the fixed 3pt default glues together.
                words = _join_touching(page.extract_words(use_text_flow=False, keep_blank_chars=False,
                                                          x_tolerance_ratio=0.2))
                pages.append({"words": words, "width": float(page.width), "text": text,
                              "source": "pdf-text"})
            else:
                pages.append(_ocr_page(page.to_image(resolution=OCR_DPI).original, engine,
                                       unit=OCR_DPI / 72))
    return pages


# ---- office files -------------------------------------------------------------------------

def _cell(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))      # 1200.0 -> "1200" so money parsing stays exact
    return str(v)


def read_grids(path: str) -> List[List[List[str]]]:
    """Every sheet (xlsx) or table (docx) as a grid of strings. docx paragraphs come back as
    a final one-column grid so text mode sees them too."""
    ext = ext_of(path)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [[[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
                    for ws in wb.worksheets]
        finally:
            wb.close()
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            raise ValueError("Word (.docx) files need the optional 'python-docx' package.")
        d = docx.Document(path)
        grids = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
        paragraphs = [[p.text] for p in d.paragraphs if p.text]
        return grids + ([paragraphs] if paragraphs else [])
    raise ValueError(f"not an office file: {path}")
