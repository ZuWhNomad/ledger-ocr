"""Export normalized rows to CSV (stdlib) and XLSX (openpyxl, with flags highlighted)."""
from __future__ import annotations
import csv
from typing import List, Dict

COLUMNS = ["date", "description", "debit", "credit", "amount", "balance", "flag", "llm_classification"]


def _flat(rows: List[Dict]) -> List[Dict]:
    out = []
    for r in rows:
        llm = r.get("llm") or {}
        out.append({
            "date": r.get("date", ""),
            "description": r.get("description", ""),
            "debit": r.get("debit", ""),
            "credit": r.get("credit", ""),
            "amount": r.get("amount", ""),
            "balance": r.get("balance", ""),
            "flag": r.get("flag") or "",
            "llm_classification": (llm.get("classification", "") if isinstance(llm, dict) else ""),
        })
    return out


def to_csv(rows: List[Dict], path: str) -> str:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(_flat(rows))
    return path


def to_xlsx(rows: List[Dict], path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    warn = PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid")
    for rec in _flat(rows):
        ws.append([rec[c] for c in COLUMNS])
        if rec["flag"]:
            for c in ws[ws.max_row]:
                c.fill = warn
    for i, col in enumerate(COLUMNS, 1):
        width = max(len(col), *(len(str(r[col])) for r in _flat(rows))) if rows else len(col)
        ws.column_dimensions[chr(64 + i)].width = min(max(width + 2, 8), 60)
    wb.save(path)
    return path
