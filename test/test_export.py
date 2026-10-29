import csv

import pytest

from ledger_ocr import export as XP
from ledger_ocr import reconcile as RC

ROWS, _ = RC.reconcile(RC.normalize_rows([
    {"date": "2026-02-01", "description": "open", "amount": "0.00", "balance": "1000.00"},
    {"date": "2026-02-02", "description": "dep", "amount": "200.00", "balance": "1300.00"},   # wrong
]))


def test_csv_has_every_column_and_the_flag(tmp_path):
    with open(XP.to_csv(ROWS, str(tmp_path / "o.csv")), newline="", encoding="utf-8-sig") as f:   # BOM: Excel opens it as UTF-8
        out = list(csv.DictReader(f))
    assert list(out[0]) == [c for c in XP.COLUMNS if c not in ("type", "num", "account")]   # absent here
    assert [bool(r["flag"]) for r in out] == [False, True]


def test_general_ledger_columns_appear_when_present(tmp_path):
    rows = [dict(ROWS[0], type="Payment", num="2998")]
    with open(XP.to_csv(rows, str(tmp_path / "g.csv")), newline="", encoding="utf-8-sig") as f:
        assert next(csv.reader(f))[:4] == ["type", "date", "num", "description"]


def test_xlsx_opens(tmp_path):
    import openpyxl
    ws = openpyxl.load_workbook(XP.to_xlsx(ROWS, str(tmp_path / "o.xlsx"))).active
    assert ws.max_row == 3


def test_docx_table_and_text(tmp_path):
    docx = pytest.importorskip("docx")
    table = docx.Document(XP.to_docx(ROWS, str(tmp_path / "t.docx"))).tables[0]
    assert len(table.rows[0].cells) == len(XP.COLUMNS) - 3     # no type/num/account in these rows
    doc = docx.Document(XP.to_docx_text("line one\nline two", str(tmp_path / "x.docx")))
    assert "line one" in "\n".join(p.text for p in doc.paragraphs)
