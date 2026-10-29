"""End to end: file in -> rows/text + output files."""
import os
from decimal import Decimal

import pytest

from conftest import FIXTURES
from ledger_ocr import readers as R
from ledger_ocr.pipeline import process

SAMPLE = os.path.join(FIXTURES, "bank_statement.pdf")
needs_engine = pytest.mark.parametrize(
    "engine", R.available_engines() or [pytest.param("none", marks=pytest.mark.skip("no OCR engine installed"))])


def _printed_sheet(path, size):
    """A spreadsheet printed to PDF, scaled to `size` pt type: gridlines drawn per cell, a note
    that overflows into the empty cell beside it (no gridline under it), accounting padding
    brackets drawn fully transparent, and the header row repeated on page 2."""
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    k = size / 8
    xs = [50 * k, 110 * k, 170 * k, 330 * k, 400 * k]
    c = canvas.Canvas(str(path))
    pages = [[("Date", "KM", "Notes"), ("Jan26", "24.00", "Quote"),
              ("Feb 5th", "67.40", "supplies run and a long note that overflows"), ("Mar 3", "$ 2.74", "Paint")],
             [("Date", "KM", "Notes"), ("Apr 1", "10.00", "Hardware store"), ("Apr 2", "11.50", "Client visit")]]
    for rows in pages:
        c.setFont("Helvetica", size)
        for n, (date, km, note) in enumerate(rows):
            top = 800 - 12 * k * n
            for x in xs:
                if not (x == xs[3] and "overflows" in note):
                    c.line(x, top, x, top - 12 * k)
            c.drawString(xs[0] + 2 * k, top - 9 * k, date)
            c.drawRightString(xs[2] - 2 * k, top - 9 * k, km)
            c.drawString(xs[2] + 2 * k, top - 9 * k, note)
            if km.startswith("$"):
                c.setFillAlpha(0)
                c.drawString(xs[1] + 1 * k, top - 9 * k, "(")
                c.drawString(xs[2] - 1.9 * k, top - 9 * k, ")")
                c.setFillAlpha(1)
        c.showPage()
    c.save()


@pytest.mark.parametrize("size", [8, 2])     # 2pt: a whole sheet shrunk onto one page
def test_printed_spreadsheet_is_cut_along_its_gridlines(tmp_path, size):
    _printed_sheet(tmp_path / "sheet.pdf", size)
    r = process(str(tmp_path / "sheet.pdf"))
    assert r["document_shape"] == "table"
    assert r["grid"] == [["Date", "KM", "Notes"], ["Jan26", "24.00", "Quote"],
                         ["Feb 5th", "67.40", "supplies run and a long note that overflows"],
                         ["Mar 3", "$ 2.74", "Paint"],          # invisible '(' ')' never read
                         ["Apr 1", "10.00", "Hardware store"],  # page 2's repeated header kept once
                         ["Apr 2", "11.50", "Client visit"]]


def test_statement_extracts_and_flags_the_one_bad_balance(tmp_path):
    r = process(SAMPLE, outdir=str(tmp_path))
    assert r["ok"] and r["document_shape"] == "ledger"
    assert r["summary"]["sources"] == ["pdf-text"]
    assert r["summary"]["rows"] == 15, r["summary"]
    assert all(row.get("description") for row in r["rows"])
    flagged = [row for row in r["rows"] if row.get("flag")]
    assert len(flagged) == 1 and flagged[0]["date"] == "2026-01-14", flagged
    assert r["summary"]["sign_convention"] == "debit_lowers_balance"
    assert all(os.path.exists(r["outputs"][k]) for k in ("csv", "xlsx"))


def test_single_amount_pdf():
    r = process(os.path.join(FIXTURES, "single_amount.pdf"))
    assert r["summary"]["sign_convention"] == "single_amount_column", r["summary"]
    assert r["summary"]["balance_mismatches"] == 1, r["summary"]


def test_two_page_no_header_keeps_all_rows():
    """A continuation page whose header does not repeat must still yield its rows."""
    r = process(os.path.join(FIXTURES, "two_page.pdf"))
    assert [row.get("date") for row in r["rows"]] == ["2026-01-02", "2026-01-03", "2026-01-04", "2026-01-05"]
    assert r["summary"]["balance_mismatches"] == 0, r["summary"]


def test_description_does_not_bleed_into_amount(tmp_path):
    """A wide description must not leak a token into the right-aligned amount column."""
    import make_sample
    path = str(tmp_path / "bleed.pdf")
    make_sample.make_bleed_pdf(path)
    check = [r for r in process(path)["rows"] if r.get("date") == "2026-01-03"][0]
    assert check["_amount"] == Decimal("100.00"), check
    assert "CHECK" in (check.get("description") or "")


def test_prose_pdf_becomes_text_not_an_empty_csv(tmp_path):
    import make_sample
    path = str(tmp_path / "prose.pdf")
    make_sample.make_prose_pdf(path)
    r = process(path, outdir=str(tmp_path))
    assert r["ok"] and r["document_shape"] == "text" and r["rows"] == []
    assert "csv" not in r["outputs"]
    with open(r["outputs"]["text"], encoding="utf-8") as f:
        assert "powerhouse of the cell" in f.read()


def test_text_mode_forces_text(tmp_path):
    r = process(SAMPLE, outdir=str(tmp_path), mode="text")
    assert r["document_shape"] == "text" and r["rows"] == []
    assert "Balance" in r["text_preview"]


@needs_engine
def test_scanned_statement_recovers_columns_and_the_flag(statement_png, tmp_path, engine):
    r = process(statement_png, outdir=str(tmp_path), engine=engine)
    assert r["ok"] and r["document_shape"] == "ledger", r
    assert r["summary"]["sources"] == [engine]
    assert any(row.get("debit") for row in r["rows"]) and any(row.get("credit") for row in r["rows"])
    assert r["summary"]["balance_mismatches"], r["summary"]
    assert os.path.exists(r["outputs"]["xlsx"])


def test_xlsx_ledger(tmp_path):
    import openpyxl
    path = str(tmp_path / "ledger.xlsx")
    wb = openpyxl.Workbook()
    for row in (["Date", "Description", "Debit", "Credit", "Balance"],
                ["2026-01-01", "Opening Balance", None, None, 1000.00],
                ["2026-01-02", "Deposit", None, 500.00, 1500.00],
                ["2026-01-03", "Deposit", None, 500.00, 2000.00],
                ["2026-01-04", "Withdrawal", 200.00, None, 1200.00]):     # should be 1800
        wb.active.append(row)
    wb.save(path)
    r = process(path)
    assert r["ok"] and r["summary"]["sources"] == ["grid"]
    assert len(r["rows"]) == 4 and r["summary"]["balance_mismatches"] >= 1, r["summary"]


def test_docx_table_and_docx_prose(tmp_path):
    docx = pytest.importorskip("docx")
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=3)
    for i, row in enumerate((["Date", "Description", "Balance"], ["2026-01-01", "Open", "1000.00"],
                             ["2026-01-02", "Same", "1000.00"])):
        for j, val in enumerate(row):
            table.cell(i, j).text = val
    doc.save(str(tmp_path / "t.docx"))
    assert len(process(str(tmp_path / "t.docx"))["rows"]) == 2

    prose = docx.Document()
    prose.add_paragraph("Only paragraph text.")
    prose.save(str(tmp_path / "p.docx"))
    r = process(str(tmp_path / "p.docx"), outdir=str(tmp_path), mode="table")
    assert r["document_shape"] == "text" and "csv" not in r["outputs"]
    assert "Only paragraph text." in r["text_preview"]


def test_errors_are_returned_not_raised(tmp_path):
    assert not process(str(tmp_path / "missing.pdf"))["ok"]
    assert not process(SAMPLE, mode="nope")["ok"]


def test_unknown_engine_on_a_scan_is_an_error(statement_png):
    assert "unknown engine" in process(statement_png, engine="nope")["error"]


PRIVATE = os.environ.get("LEDGEROCR_PRIVATE_FIXTURES", "")


@pytest.mark.skipif(not os.path.isdir(PRIVATE), reason="LEDGEROCR_PRIVATE_FIXTURES not set")
@pytest.mark.parametrize("name", sorted(f for f in os.listdir(PRIVATE) if "." in f and ".expected." not in f)
                         if os.path.isdir(PRIVATE) else [])
def test_private_real_documents_are_read(name, tmp_path):
    """Smoke only (accuracy is bench/'s job): every real document yields rows or real text."""
    r = process(os.path.join(PRIVATE, name), outdir=str(tmp_path))
    assert r["ok"], r
    assert r["rows"] or len(r["text_preview"].split()) > 20, r
