"""Export normalized rows to CSV (stdlib) and XLSX (openpyxl, with flags highlighted)."""
from __future__ import annotations
import csv
from decimal import Decimal
from typing import List, Dict

COLUMNS = ["date", "description", "debit", "credit", "amount", "balance", "flag", "llm_classification"]
# Money columns and the parsed-Decimal field that backs each (see reconcile.normalize_rows).
_MONEY_FIELD = {"debit": "_debit", "credit": "_credit", "amount": "_amount", "balance": "_balance"}


def _flat(rows: List[Dict]) -> List[Dict]:
    out = []
    for r in rows:
        llm = r.get("llm") or {}
        out.append({
            "date": r.get("date", ""),
            "description": r.get("description", ""),
            "debit": r.get("debit", ""),
            "credit": r.get("credit", ""),
            "amount": r.get("amount", ""),
            "balance": r.get("balance", ""),
            "flag": r.get("flag") or "",
            "llm_classification": (llm.get("classification", "") if isinstance(llm, dict) else ""),
        })
    return out


def to_csv(rows: List[Dict], path: str) -> str:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(_flat(rows))
    return path


def _xlsx_cell(row: Dict, col: str):
    """Cell value for XLSX: money columns become real numbers (exact parsed Decimal, so
    Excel can sum/filter them and no float rounding is introduced); everything else is the
    string projection. An unparseable money cell falls back to its original text."""
    field = _MONEY_FIELD.get(col)
    if field is not None:
        val = row.get(field)
        if isinstance(val, Decimal):
            return val  # openpyxl writes Decimal as a numeric cell, value-exact
    return None  # signal: use the flat string


def to_xlsx(rows: List[Dict], path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    warn = PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid")
    flat = _flat(rows)
    for row, frec in zip(rows, flat):
        values = []
        for col in COLUMNS:
            num = _xlsx_cell(row, col)
            values.append(num if num is not None else frec[col])
        ws.append(values)
        if frec["flag"]:
            for c in ws[ws.max_row]:
                c.fill = warn
    for i, col in enumerate(COLUMNS, 1):
        width = max(len(col), *(len(str(r[col])) for r in flat)) if flat else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    wb.save(path)
    return path
