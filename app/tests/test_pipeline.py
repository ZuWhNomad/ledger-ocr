"""End-to-end tests + the A/B extraction-strategy comparison.

Runnable two ways:
    python -m pytest tests/            # if pytest is installed
    python tests/test_pipeline.py      # plain runner, no pytest needed
"""
from __future__ import annotations
import os
import sys
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from ocr_pipeline.pipeline import process, parse_text_rows  # noqa: E402
from ocr_pipeline import reconcile as RC  # noqa: E402
from ocr_pipeline import extract as EX  # noqa: E402

SAMPLE = os.path.join(ROOT, "samples", "bank_statement.pdf")


def _ensure_sample():
    if not os.path.exists(SAMPLE):
        sys.path.insert(0, os.path.join(ROOT, "samples"))
        import make_sample
        make_sample.make_pdf(SAMPLE)


def _make_fixture(builder, name):
    """Build (once) a fixture PDF under samples/ using the named make_sample builder."""
    sys.path.insert(0, os.path.join(ROOT, "samples"))
    import make_sample
    path = os.path.join(ROOT, "samples", name)
    getattr(make_sample, builder)(path)  # deterministic; cheap to regenerate
    return path


def test_words_strategy_extracts_and_flags():
    _ensure_sample()
    r = process(SAMPLE, strategy="words")
    assert r["ok"] and r["route"] == "born-digital"
    assert r["summary"]["rows"] == 15, r["summary"]
    # every row has a non-empty description (field fidelity)
    assert all(row.get("description") for row in r["rows"])
    # the one injected error is flagged, on the right row, and nothing else
    flagged = [row for row in r["rows"] if row.get("flag")]
    assert len(flagged) == 1, flagged
    assert flagged[0]["date"] == "2026-01-14"
    assert r["summary"]["sign_convention"] == "debit_lowers_balance"


def test_ab_words_beats_lines_on_fidelity():
    """A/B: both may capture rows, but 'words' preserves the description column."""
    _ensure_sample()
    def desc_fidelity(strat):
        rows = process(SAMPLE, strategy=strat)["rows"]
        if not rows:
            return 0.0
        return sum(1 for x in rows if (x.get("description") or "").strip()) / len(rows)
    words = desc_fidelity("words")
    lines = desc_fidelity("lines")
    print(f"\n[A/B] description fidelity: words={words:.2f} lines={lines:.2f} -> "
          f"winner={'words' if words >= lines else 'lines'}")
    assert words == 1.0
    assert words >= lines  # header-anchored clustering never loses to generic text tables


def test_parse_money_edge_cases():
    p = RC.parse_money
    assert p("1,234.56") == Decimal("1234.56")
    assert p("$2,000.00") == Decimal("2000.00")
    assert p("(84.19)") == Decimal("-84.19")   # parentheses = negative
    assert p("-25.00") == Decimal("-25.00")
    assert p("1.234,56") == Decimal("1234.56")     # EU thousands/decimal
    assert p("1 234,56") == Decimal("1234.56")     # EU with space thousands
    assert p("84.19 DR") == Decimal("-84.19")      # trailing debit marker
    assert p("1,234.56 CR") == Decimal("1234.56")  # credit marker stays positive
    assert p("1234.56-") == Decimal("-1234.56")    # trailing minus
    assert p("−1,234.56") == Decimal("-1234.56")  # unicode minus
    assert p("") is None
    assert p("N/A") is None
    assert p("OCR-Pipeline tool.") is None


def test_sign_autodetect_ledger_debit_raises():
    """Asset-side ledger where debits INCREASE the balance -> auto-detected."""
    raw = [
        {"date": "2026-02-01", "description": "open", "balance": "100.00"},
        {"date": "2026-02-02", "description": "buy", "debit": "40.00", "balance": "140.00"},
        {"date": "2026-02-03", "description": "sell", "credit": "10.00", "balance": "130.00"},
    ]
    rows = RC.normalize_rows(raw)
    rows, summary = RC.reconcile(rows)
    assert summary["sign_convention"] == "debit_raises_balance", summary
    assert summary["balance_mismatches"] == 0


def test_parse_money_locale_thousands():
    """COR-P1-1: US thousands, EU thousands, whole-dollar and millions parse exactly.

    Regression for the old 'locale = whichever separator is last' guess, which turned
    '5,000' into 5.00 and dropped '1.000.000' entirely.
    """
    p = RC.parse_money
    # US thousands, whole dollars (no decimals) -- the core silent bug
    assert p("5,000") == Decimal("5000")
    assert p("1,500") == Decimal("1500")
    assert p("12,345") == Decimal("12345")
    assert p("1,234,567") == Decimal("1234567")
    # US thousands with cents
    assert p("1,234.56") == Decimal("1234.56")
    assert p("12,345.67") == Decimal("12345.67")
    # EU thousands (dot groups), with and without decimals
    assert p("1.234") == Decimal("1234")
    assert p("1.000.000") == Decimal("1000000")
    assert p("1.234,56") == Decimal("1234.56")
    assert p("1 234,56") == Decimal("1234.56")      # EU space thousands
    assert p("1 000 000") == Decimal("1000000")
    # whole-dollar plain integer and a plain cents value
    assert p("5000") == Decimal("5000")
    assert p("100.00") == Decimal("100.00")
    # a money cell that is really two tokens (leaked description) must be REJECTED,
    # never glued into one giant number
    assert p("1043 100.00") is None
    assert p("1,234.56 (2)") is None


def test_single_amount_column_reconciles():
    """COR-P1-2 / P2-1: a single signed-amount column must be reconciled, not declared
    'unknown'. A running-balance typo has to be flagged (was silently passing)."""
    raw = [
        {"date": "2026-02-01", "description": "open", "amount": "0.00", "balance": "1000.00"},
        {"date": "2026-02-02", "description": "dep", "amount": "200.00", "balance": "1200.00"},
        {"date": "2026-02-03", "description": "wire", "amount": "-350.00", "balance": "999.00"},  # should be 850
        {"date": "2026-02-04", "description": "refund", "amount": "50.00", "balance": "900.00"},
    ]
    rows = RC.normalize_rows(raw)
    rows, summary = RC.reconcile(rows)
    assert summary["sign_convention"] == "single_amount_column", summary
    assert summary["balance_checked"] is True, summary
    assert summary["balance_mismatches"] == 1, summary
    flagged = [r for r in rows if r.get("flag")]
    assert len(flagged) == 1 and flagged[0]["date"] == "2026-02-03", flagged


def test_no_balance_reports_not_checked():
    """COR-P2-8: a statement with no balance column must report 'not checked', not a
    misleading '0 mismatches'."""
    raw = [{"date": "2026-02-01", "description": "a", "amount": "10.00"},
           {"date": "2026-02-02", "description": "b", "amount": "-30.00"}]
    rows, summary = RC.reconcile(RC.normalize_rows(raw))
    assert summary["balance_checked"] is False, summary
    assert summary["balance_mismatches"] is None, summary


def test_single_amount_pdf_end_to_end():
    """Same as above but through the full extract->reconcile pipeline from a PDF fixture."""
    path = _make_fixture("make_single_amount_pdf", "fixture_single_amount.pdf")
    r = process(path, strategy="words")
    assert r["ok"], r
    assert r["summary"]["sign_convention"] == "single_amount_column", r["summary"]
    assert r["summary"]["balance_mismatches"] == 1, r["summary"]
    flagged = [row for row in r["rows"] if row.get("flag")]
    assert len(flagged) == 1 and flagged[0]["date"] == "2026-02-03", flagged


def test_two_page_no_header_keeps_all_rows():
    """COR-P2-5: a continuation page whose header does not repeat must still yield its rows
    (they used to vanish silently)."""
    path = _make_fixture("make_two_page_pdf", "fixture_two_page.pdf")
    r = process(path, strategy="words")
    assert r["ok"], r
    dates = [row.get("date") for row in r["rows"]]
    assert dates == ["2026-01-02", "2026-01-03", "2026-01-04", "2026-01-05"], dates
    assert r["summary"]["balance_mismatches"] == 0, r["summary"]  # all four reconcile


def test_description_does_not_bleed_into_amount():
    """COR-P1-3: a wide description must not leak a token into the right-aligned amount
    column and concatenate into a giant fake number. Real amount is 100.00."""
    path = _make_fixture("make_bleed_pdf", "fixture_bleed.pdf")
    rows = process(path, strategy="words")["rows"]
    check = [r for r in rows if r.get("date") == "2026-01-03"][0]
    assert check["_amount"] == Decimal("100.00"), check
    assert "CHECK" in (check.get("description") or "")  # leaked tokens landed back in description


def test_ocr_textline_parser():
    """The scanned-page fallback parser reads date + trailing amounts from raw text."""
    text = "2026-03-01 COFFEE SHOP 12.50 987.50\n2026-03-02 PAYCHECK 2,000.00 2,987.50"
    rows = parse_text_rows(text)
    assert len(rows) == 2
    assert rows[0]["date"] == "2026-03-01"
    assert rows[0]["balance"] == "987.50"
    assert rows[1]["amount"] == "2,000.00"


def test_ocr_words_path_recovers_columns():
    """Scanned path: the image_to_data word-box reconstruction recovers separate debit/credit
    columns (the plain-text parser cannot) and still flags the injected balance error.
    Skips cleanly if the Tesseract OCR engine is not installed."""
    from ocr_pipeline import extract as EX
    if not EX.ocr_available():
        print("SKIP test_ocr_words_path_recovers_columns (Tesseract not installed)")
        return
    import tempfile
    import pdfplumber
    _ensure_sample()
    with tempfile.TemporaryDirectory() as td:
        png = os.path.join(td, "scan.png")
        with pdfplumber.open(SAMPLE) as pdf:
            pdf.pages[0].to_image(resolution=300).original.save(png)
        r = process(png, strategy="words")
    assert r["ok"], r
    assert r["route"] == "ocr", r["route"]
    assert r["summary"]["strategy"] == "ocr-words", r["summary"]
    # debit AND credit recovered as distinct columns (impossible via image_to_string path)
    assert any(row.get("debit") for row in r["rows"]), r["rows"]
    assert any(row.get("credit") for row in r["rows"]), r["rows"]
    # the deliberately corrupted running balance is still caught
    assert r["summary"]["balance_mismatches"], r["summary"]


def test_xlsx_input_end_to_end():
    import tempfile
    import openpyxl
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "test_ledger.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Date", "Description", "Debit", "Credit", "Balance"])
        ws.append(["2026-01-01", "Opening Balance", None, None, 1000.00])
        ws.append(["2026-01-02", "Deposit", None, 500.00, 1500.00])
        ws.append(["2026-01-03", "Deposit", None, 500.00, 2000.00])
        ws.append(["2026-01-04", "Withdrawal", 200.00, None, 1200.00])
        wb.save(path)
        wb.close()

        r = process(path)
        assert r["ok"], r
        assert r["route"] == "structured", r["route"]
        assert len(r["rows"]) == 4, r["rows"]
        assert r["summary"]["balance_mismatches"] >= 1, r["summary"]


def test_docx_input_end_to_end():
    try:
        import docx
    except ImportError:
        print("SKIP test_docx_input_end_to_end (python-docx not installed)")
        return
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "test_ledger.docx")
        doc = docx.Document()
        table = doc.add_table(rows=4, cols=5)
        headers = ["Date", "Description", "Debit", "Credit", "Balance"]
        for j, h in enumerate(headers):
            table.cell(0, j).text = h
        data = [
            ["2026-01-01", "Opening Balance", "", "", "1000.00"],
            ["2026-01-02", "Deposit", "", "500.00", "1500.00"],
            ["2026-01-03", "Withdrawal", "200.00", "", "1300.00"],
        ]
        for i, row in enumerate(data, start=1):
            for j, val in enumerate(row):
                table.cell(i, j).text = val
        doc.save(path)

        r = process(path)
        assert r["ok"], r
        assert r["route"] == "structured", r["route"]
        assert len(r["rows"]) == 3, r["rows"]


def test_docx_export():
    try:
        import docx
    except ImportError:
        print("SKIP test_docx_export (python-docx not installed)")
        return
    import tempfile
    _ensure_sample()
    with tempfile.TemporaryDirectory() as td:
        r = process(SAMPLE, outdir=td)
        assert r["ok"], r
        assert "docx" in r["outputs"], r["outputs"]
        docx_path = r["outputs"]["docx"]
        assert os.path.exists(docx_path)
        doc = docx.Document(docx_path)
        assert len(doc.tables) >= 1
        tbl = doc.tables[0]
        from ocr_pipeline.export import COLUMNS
        assert len(tbl.rows[0].cells) == len(COLUMNS)


def test_get_model_precedence():
    import tempfile
    from ocr_pipeline import validate as VAL
    old_env = os.environ.get("OCR_LLM_MODEL")
    old_cfg_env = os.environ.get("LEDGEROCR_CONFIG")
    with tempfile.TemporaryDirectory() as td:
        cfg_file = os.path.join(td, "config.json")
        os.environ["LEDGEROCR_CONFIG"] = cfg_file
        VAL.set_model("config-model")
        os.environ["OCR_LLM_MODEL"] = "env-model"
        try:
            assert VAL.get_model() == "env-model"
        finally:
            if old_env is not None:
                os.environ["OCR_LLM_MODEL"] = old_env
            else:
                os.environ.pop("OCR_LLM_MODEL", None)
            if old_cfg_env is not None:
                os.environ["LEDGEROCR_CONFIG"] = old_cfg_env
            else:
                os.environ.pop("LEDGEROCR_CONFIG", None)


def test_set_get_model_roundtrip():
    import tempfile
    from ocr_pipeline import validate as VAL
    old_env = os.environ.pop("OCR_LLM_MODEL", None)
    old_cfg_env = os.environ.get("LEDGEROCR_CONFIG")
    with tempfile.TemporaryDirectory() as td:
        cfg_file = os.path.join(td, "config.json")
        os.environ["LEDGEROCR_CONFIG"] = cfg_file
        try:
            VAL.set_model("custom-model-x")
            assert VAL.get_model() == "custom-model-x"
        finally:
            if old_env is not None:
                os.environ["OCR_LLM_MODEL"] = old_env
            if old_cfg_env is not None:
                os.environ["LEDGEROCR_CONFIG"] = old_cfg_env
            else:
                os.environ.pop("LEDGEROCR_CONFIG", None)


def test_suggest_from_prefers_default_then_smallest():
    from ocr_pipeline import validate as VAL
    # 1. Include DEFAULT_MODEL fitting+json_ok, a big one not fitting, a fitting one that failed json
    models1 = [
        {"name": VAL.DEFAULT_MODEL, "size_gb": 1.9, "fits_ram": True, "json_ok": True, "latency_s": 3.6},
        {"name": "huge-model", "size_gb": 30.0, "fits_ram": False, "json_ok": None, "latency_s": None},
        {"name": "broken-model", "size_gb": 1.0, "fits_ram": True, "json_ok": False, "latency_s": 1.2},
    ]
    assert VAL.suggest_from(models1) == VAL.DEFAULT_MODEL

    # 2. List WITHOUT DEFAULT_MODEL where two fit and pass -> returns the smaller
    models2 = [
        {"name": "model-b", "size_gb": 5.0, "fits_ram": True, "json_ok": True, "latency_s": 2.0},
        {"name": "model-a", "size_gb": 2.0, "fits_ram": True, "json_ok": True, "latency_s": 4.0},
    ]
    assert VAL.suggest_from(models2) == "model-a"

    # 3. List where none pass json but some fit -> returns smallest fitting
    models3 = [
        {"name": "fit-large", "size_gb": 8.0, "fits_ram": True, "json_ok": False, "latency_s": 1.0},
        {"name": "fit-small", "size_gb": 3.0, "fits_ram": True, "json_ok": False, "latency_s": 1.0},
        {"name": "no-fit", "size_gb": 25.0, "fits_ram": False, "json_ok": None, "latency_s": None},
    ]
    assert VAL.suggest_from(models3) == "fit-small"

    # 4. Empty list -> None
    assert VAL.suggest_from([]) is None


def test_webp_and_heic_are_images():
    assert ".webp" in EX.IMAGE_EXTS and ".heic" in EX.IMAGE_EXTS and EX.is_image("x.webp")


def test_docx_no_table_extracts_text():
    try:
        import docx
    except ImportError:
        print("SKIP test_docx_no_table_extracts_text (python-docx not installed)")
        return
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "no_table.docx")
        doc = docx.Document()
        doc.add_paragraph("This is a document with text but no tables.")
        doc.save(path)

        r = process(path, outdir=td)
        assert r["ok"], r
        assert r["document_shape"] == "text", r["document_shape"]
        assert len(r["rows"]) == 0, r["rows"]
        assert "csv" not in r["outputs"], r["outputs"]
        assert r["outputs"].get("text") and os.path.exists(r["outputs"]["text"])
        assert r["outputs"].get("docx")
        with open(r["outputs"]["text"], "r", encoding="utf-8") as f:
            content = f.read()
        assert "document" in content and "tables" in content


def test_text_mode_forces_text():
    _ensure_sample()
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        r = process(SAMPLE, outdir=td, mode="text")
        assert r["document_shape"] == "text", r
        assert r["outputs"].get("text")
        with open(r["outputs"]["text"], "r", encoding="utf-8") as f:
            assert f.read().strip()
        assert r["rows"] == []


def test_table_mode_zero_rows_falls_back_to_text():
    try:
        import docx
    except ImportError:
        print("SKIP test_table_mode_zero_rows_falls_back_to_text (python-docx not installed)")
        return
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        path = os.path.join(td, "no_table.docx")
        doc = docx.Document()
        doc.add_paragraph("Only paragraph text.")
        doc.save(path)
        r = process(path, outdir=td, mode="table")
        assert r["document_shape"] == "text", r
        assert "csv" not in r["outputs"]
        assert r["outputs"].get("text")


def test_ledger_still_recognized():
    _ensure_sample()
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        r = process(SAMPLE, outdir=td, strategy="words")
        assert r["ok"], r
        assert r["document_shape"] == "ledger", r["document_shape"]
        assert len(r["rows"]) > 0, r["rows"]
        assert r["outputs"].get("csv") and os.path.exists(r["outputs"]["csv"])


def test_receipt_photo_extracts_text():
    jpg_path = os.path.join(ROOT, "tests", "fixtures", "real", "receipt_pharmacy.jpg")
    if not EX.ocr_available():
        print("SKIP test_receipt_photo_extracts_text (Tesseract not installed)")
        return
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        r = process(jpg_path, outdir=td)
        assert r["ok"], r
        assert r["document_shape"] == "text", r["document_shape"]
        assert len(r["rows"]) == 0, r["rows"]
        assert "csv" not in r["outputs"], r["outputs"]
        text_file = r["outputs"].get("text")
        assert text_file and os.path.exists(text_file)
        with open(text_file, "r", encoding="utf-8") as f:
            content = f.read()
        assert len(content) > 0, content


def _run():
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    failed = 0
    for t in tests:
        try:
            t()
            print(f"PASS {t.__name__}")
        except AssertionError as e:
            failed += 1
            print(f"FAIL {t.__name__}: {e}")
        except Exception as e:
            failed += 1
            print(f"ERROR {t.__name__}: {type(e).__name__}: {e}")
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    return failed


if __name__ == "__main__":
    sys.exit(1 if _run() else 0)

